import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("Installing python-pptx...")
    os.system("pip install python-pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE


def create_genericrx_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Theme Colors
    DARK_BG = RGBColor(15, 23, 42)      # Deep Slate #0f172a
    CARD_BG = RGBColor(30, 41, 59)      # Slate #1e293b
    TEAL = RGBColor(20, 184, 166)       # Teal #14b8a6
    EMERALD = RGBColor(16, 185, 129)    # Emerald #10b981
    WHITE = RGBColor(248, 250, 252)     # White #f8fafc
    MUTED = RGBColor(148, 163, 184)     # Muted #94a3b8
    CYAN = RGBColor(6, 182, 212)        # Cyan #06b6d4

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="GENERICRX COLLEGE MINI-PROJECT"):
        # Header text container
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = TEAL
        p1.font.name = "Arial"

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(26)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.font.name = "Arial"

    # Slide 1: Title Slide
    slide1 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide1)

    # Title Card Accent Box
    card1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = CARD_BG
    card1.line.color.rgb = TEAL
    card1.line.width = Pt(2)

    tb1 = slide1.shapes.add_textbox(Inches(2.0), Inches(2.0), Inches(9.333), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "GenericRx"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = TEAL
    p.font.name = "Arial"

    p = tf1.add_paragraph()
    p.text = "Offline-First Medicine Alternative & Cost Audit System"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Arial"

    p = tf1.add_paragraph()
    p.text = "\nDomain: Healthcare Tech, Progressive Web Applications (PWA), Data Engineering"
    p.font.size = Pt(14)
    p.font.color.rgb = MUTED

    p = tf1.add_paragraph()
    p.text = "Technology Stack: Python 3.14 (FastAPI + SQLite), Leaflet.js, PWA (Service Worker + IndexedDB)"
    p.font.size = Pt(14)
    p.font.color.rgb = CYAN

    p = tf1.add_paragraph()
    p.text = "\nPresented by: College Mini-Project Team"
    p.font.size = Pt(13)
    p.font.color.rgb = EMERALD

    # Slide 2: Problem Statement & Background
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide2)
    add_header(slide2, "1. Problem Statement & Real-World Background")

    box_data = [
        ("💊 High Cost of Brand Medications", "Patients frequently visit pharmacies near hospitals or rural clinics and find prescribed brand-name drugs out of stock or unaffordable.\n\nLack of awareness about government-approved generic alternatives with the exact same Active Pharmaceutical Ingredient (API)."),
        ("📵 Zero-Connectivity Blind Spots", "Hospital basements, rural clinics, and emergency rooms frequently have weak or zero cellular signal (2G/3G/blind spots).\n\nStandard online health apps completely crash or fail to load when patients need them most."),
        ("⚠️ Patient Safety & Regulatory Risks", "Substitutions without checking active salt strength or dosage form (Tablet vs Syrup) cause medical risks.\n\nNeed for clear prescription regulatory tags (Schedule H / Schedule X).")
    ]

    lefts = [0.8, 4.8, 8.8]
    for i, (title, content) in enumerate(box_data):
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(lefts[i]), Inches(1.8), Inches(3.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = TEAL if i == 0 else MUTED

        tb = slide2.shapes.add_textbox(Inches(lefts[i] + 0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = TEAL if i == 0 else WHITE

        p2 = tf.add_paragraph()
        p2.text = "\n" + content
        p2.font.size = Pt(13)
        p2.font.color.rgb = MUTED

    # Slide 3: Objectives & Solution Overview
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide3)
    add_header(slide3, "2. Objectives & Core Functional Modules")

    modules = [
        ("Module 1: Intelligent Search", "Instant Brand-to-Salt mapping (e.g. Crocin 650 -> Paracetamol 650mg) with typo-tolerant fuzzy search."),
        ("Module 2: Cost Audit Engine", "Direct price savings formula calculation, visual color badges, and monthly/annual prescription savings audit."),
        ("Module 3: Offline-First PWA", "Pre-seeded IndexedDB + Service Worker caching for 100% offline functionality without internet connection."),
        ("Module 4: Safety & Regulatory", "Schedule H & Schedule X prescription alerts plus dosage strength/form mismatch warnings."),
        ("Module 5: 100% Free Locator", "Interactive Leaflet.js & OpenStreetMap locator for nearby Jan Aushadhi Kendras & retail pharmacies.")
    ]

    top_pos = 1.8
    for title, desc in modules:
        card = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_pos), Inches(11.733), Inches(0.9))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = TEAL

        tb = slide3.shapes.add_textbox(Inches(1.0), Inches(top_pos + 0.1), Inches(11.3), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title + ": "
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = EMERALD

        run = p.add_run()
        run.text = desc
        run.font.size = Pt(14)
        run.font.bold = False
        run.font.color.rgb = WHITE

        top_pos += 1.05

    # Slide 4: System Architecture
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide4)
    add_header(slide4, "3. System Architecture & Offline Cache Flow")

    arch_box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    arch_box.fill.solid()
    arch_box.fill.fore_color.rgb = CARD_BG
    arch_box.line.color.rgb = CYAN

    tb = slide4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Dual-Tier Architecture with Seamless Network Failover:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL

    arch_text = (
        "1. ONLINE MODE:\n"
        "   • Client PWA sends requests to Python FastAPI REST API Backend.\n"
        "   • Queries SQLite database (genericrx.db) containing 567 seeded medicines and store locations.\n"
        "   • Python RapidFuzz engine performs server-side fuzzy matching.\n\n"
        "2. OFFLINE SYNC LAYER:\n"
        "   • Service Worker (sw.js) caches static app shell (HTML/CSS/JS).\n"
        "   • Client-side IndexedDB (Dexie.js) pre-syncs the dataset from Python server.\n\n"
        "3. OFFLINE MODE (ZERO CONNECTIVITY):\n"
        "   • When cell network fails, app seamlessly switches to IndexedDB + Fuse.js fuzzy engine.\n"
        "   • Patients retain 100% search, savings calculation, and map lookup in hospital basements."
    )
    p2 = tf.add_paragraph()
    p2.text = arch_text
    p2.font.size = Pt(14)
    p2.font.color.rgb = WHITE

    # Slide 5: Database & Entity-Relationship Schema
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide5)
    add_header(slide5, "4. Database Schema & Data Models")

    # Table 1: Medicine
    t1 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    t1.fill.solid()
    t1.fill.fore_color.rgb = CARD_BG
    t1.line.color.rgb = TEAL

    tb1 = slide5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(5.2), Inches(4.4))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "Medicine Table Schema"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL

    schema1_text = (
        "• id: INTEGER (Primary Key)\n"
        "• name: VARCHAR(255) (e.g. Crocin 650)\n"
        "• salt_name: VARCHAR(255) (Paracetamol 650mg)\n"
        "• strength: VARCHAR(100) (650mg)\n"
        "• dosage_form: VARCHAR(100) (Tablet / Syrup)\n"
        "• price_per_unit: FLOAT (Price in INR)\n"
        "• is_generic: BOOLEAN (Jan Aushadhi Flag)\n"
        "• manufacturer: VARCHAR(255)\n"
        "• schedule: VARCHAR(50) (OTC, Schedule H/X)\n"
        "• therapeutic_category: VARCHAR(100)"
    )
    p2 = tf1.add_paragraph()
    p2.text = schema1_text
    p2.font.size = Pt(13)
    p2.font.color.rgb = WHITE

    # Table 2: Pharmacy
    t2 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.8), Inches(5.7), Inches(4.8))
    t2.fill.solid()
    t2.fill.fore_color.rgb = CARD_BG
    t2.line.color.rgb = CYAN

    tb2 = slide5.shapes.add_textbox(Inches(7.0), Inches(2.0), Inches(5.3), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p = tf2.paragraphs[0]
    p.text = "Pharmacy Location Schema"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN

    schema2_text = (
        "• id: INTEGER (Primary Key)\n"
        "• name: VARCHAR(255) (Store Name)\n"
        "• pharmacy_type: VARCHAR(100) (Jan Aushadhi / Retail)\n"
        "• address: VARCHAR(500) (Street / Landmark)\n"
        "• city: VARCHAR(100) (Mumbai, Delhi, Bangalore)\n"
        "• phone: VARCHAR(50) (Contact Number)\n"
        "• lat, lng: FLOAT (GPS Coordinates)\n"
        "• is_open_24h: BOOLEAN (Emergency Status)\n"
        "• stock_status: VARCHAR(255) (Generics Availability)"
    )
    p2 = tf2.add_paragraph()
    p2.text = schema2_text
    p2.font.size = Pt(13)
    p2.font.color.rgb = WHITE

    # Slide 6: Mathematical Model & Cost Audit Engine
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide6)
    add_header(slide6, "5. Cost Audit Engine & Mathematical Model")

    c = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = EMERALD

    tb = slide6.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Mathematical Model for Price Comparison:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    math_text = (
        "1. Direct Percentage Cost Savings Formula:\n"
        "   Savings % = [ (Brand Price - Generic Price) / Brand Price ] * 100\n\n"
        "2. Visual Audit Indicator Badges:\n"
        "   • GREEN BADGE  : Savings >= 50% (High Cost Relief)\n"
        "   • AMBER BADGE  : Savings 20% - 50% (Moderate Savings)\n"
        "   • BLUE BADGE   : Savings < 20% (Low Savings)\n\n"
        "3. Prescription Basket Accumulator:\n"
        "   Monthly Savings = SUM [ (Brand_Price_i - Generic_Price_i) * Units_i ]\n"
        "   Annual Savings  = Monthly Savings * 12\n\n"
        "4. Printable Audit Report:\n"
        "   Generates clear patient/doctor summary with net yearly savings calculation."
    )
    p2 = tf.add_paragraph()
    p2.text = math_text
    p2.font.size = Pt(14)
    p2.font.color.rgb = WHITE

    # Slide 7: 100% Free Leaflet.js Map Locator
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide7)
    add_header(slide7, "6. 100% Free Open-Source Pharmacy Locator")

    c = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = CYAN

    tb = slide7.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Leaflet.js + OpenStreetMap Architecture (Zero API Key Cost):"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN

    map_text = (
        "• 100% Free & Open Source: No API keys, no credit cards, no billing limits.\n"
        "• Color-Coded Map Markers:\n"
        "   - GREEN PINS : Government-subsidized Jan Aushadhi Kendras (80-85% cheaper).\n"
        "   - BLUE PINS  : 24/7 Retail Pharmacies (Apollo, MedPlus, Wellness Forever).\n"
        "• GPS 'Find Near Me': Uses HTML5 navigator.geolocation for instant positioning.\n"
        "• One-Click Directions: Direct Google Maps navigation links for driving/walking.\n"
        "• Offline Backup: Store addresses & phone numbers pre-cached in IndexedDB."
    )
    p2 = tf.add_paragraph()
    p2.text = map_text
    p2.font.size = Pt(14)
    p2.font.color.rgb = WHITE

    # Slide 8: Safety & Regulatory Module
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide8)
    add_header(slide8, "7. Safety & Regulatory Warnings")

    c = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = RGBColor(244, 63, 94) # Red/Rose

    tb = slide8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Regulatory Badges & Medical Safety Flags:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = RGBColor(244, 63, 94)

    safety_text = (
        "1. Controlled Substance Warning Badges:\n"
        "   • Schedule H / H1: Prominent red warning tag ('Doctor Prescription Required').\n"
        "   • Schedule X: Animated warning tag for high-risk narcotics/stimulants.\n"
        "   • OTC: Blue tag for Over-The-Counter safe medications.\n\n"
        "2. Dosage & Form Mismatch Alert:\n"
        "   • Explicit warning when alternative differs in strength (e.g. 500mg vs 650mg).\n"
        "   • Flags dosage form mismatches (e.g. Tablet vs Syrup formulation).\n\n"
        "3. Therapeutic Safety Disclaimer:\n"
        "   • Non-prescriptive advisory urging consultation with a qualified medical practitioner."
    )
    p2 = tf.add_paragraph()
    p2.text = safety_text
    p2.font.size = Pt(14)
    p2.font.color.rgb = WHITE

    # Slide 9: Implementation & Technology Stack
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide9)
    add_header(slide9, "8. Implementation Details & Tech Stack")

    tech_boxes = [
        ("Backend Framework", "Python 3.14 + FastAPI\nUvicorn ASGI Web Server\nSQLAlchemy ORM + SQLite"),
        ("Search Engine", "Python RapidFuzz (Online)\nFuse.js Client (Offline)\nBrand-to-Salt Mapping"),
        ("Offline & Storage", "Service Worker (sw.js)\nDexie.js (IndexedDB)\nLocal Storage Cache Shell"),
        ("Frontend UI / Map", "HTML5 + Glassmorphism CSS\nLeaflet.js + OpenStreetMap\nLucide React Icons")
    ]

    positions = [(0.8, 1.8), (6.8, 1.8), (0.8, 4.4), (6.8, 4.4)]
    for i, (title, content) in enumerate(tech_boxes):
        x, y = positions[i]
        card = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.7), Inches(2.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = TEAL

        tb = slide9.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.2), Inches(5.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = TEAL

        p2 = tf.add_paragraph()
        p2.text = content
        p2.font.size = Pt(13)
        p2.font.color.rgb = WHITE

    # Slide 10: Live Demonstration Walkthrough
    slide10 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide10)
    add_header(slide10, "9. Live Demonstration & Verification")

    c = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = EMERALD

    tb = slide10.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Key Demonstration Highlights:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = EMERALD

    demo_text = (
        "1. Search Resolution Demo:\n"
        "   Search 'Crocin' -> Instantly maps to 'Paracetamol 650mg' and displays Jan Aushadhi alternatives.\n\n"
        "2. Cost Savings Audit:\n"
        "   Crocin 650 (Rs 34.50) vs Jan Aushadhi Paracetamol (Rs 5.50) -> Shows 84.1% Green Savings Badge.\n\n"
        "3. Prescription Basket & Report Export:\n"
        "   Add monthly prescription -> Generates net annual savings statement (e.g. Saved Rs 3,480 / year).\n\n"
        "4. Offline Mode Verification:\n"
        "   Disconnect internet -> Status pill changes to 'Offline Mode (IndexedDB Active)', search & map continue working."
    )
    p2 = tf.add_paragraph()
    p2.text = demo_text
    p2.font.size = Pt(14)
    p2.font.color.rgb = WHITE

    # Slide 11: Conclusion & Future Scope
    slide11 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide11)
    add_header(slide11, "10. Conclusion & Future Enhancements")

    c = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.8))
    c.fill.solid()
    c.fill.fore_color.rgb = CARD_BG
    c.line.color.rgb = TEAL

    tb = slide11.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Project Summary & Future Roadmap:"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = TEAL

    conc_text = (
        "Conclusion:\n"
        "GenericRx successfully delivers a lightweight, neutral, offline-first application that empowers patients to save up to 85% on prescription drugs while functioning reliably in zero-connectivity hospital rooms.\n\n"
        "Future Enhancements:\n"
        "• OCR Prescription Camera Scanner: Automatic extraction of brand names from paper prescriptions.\n"
        "• Multilingual Voice Search: Support for Hindi, Marathi, Tamil, and regional languages.\n"
        "• Government API Integration: Direct real-time inventory sync with National Jan Aushadhi Kendras."
    )
    p2 = tf.add_paragraph()
    p2.text = conc_text
    p2.font.size = Pt(14)
    p2.font.color.rgb = WHITE

    # Slide 12: Q&A / Final Slide
    slide12 = prs.slides.add_slide(blank_slide_layout)
    add_background(slide12)

    card12 = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2.0), Inches(1.8), Inches(9.333), Inches(4.0))
    card12.fill.solid()
    card12.fill.fore_color.rgb = CARD_BG
    card12.line.color.rgb = TEAL

    tb12 = slide12.shapes.add_textbox(Inches(2.5), Inches(2.5), Inches(8.333), Inches(3.0))
    tf12 = tb12.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "Thank You!"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = TEAL

    p = tf12.add_paragraph()
    p.text = "Questions & Answers"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = WHITE

    p = tf12.add_paragraph()
    p.text = "\nGenericRx - Offline-First Medicine Alternative & Cost Audit System"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(14)
    p.font.color.rgb = MUTED

    output_path = os.path.join(os.path.dirname(__file__), "GenericRx_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")
    return output_path


if __name__ == "__main__":
    create_genericrx_presentation()
